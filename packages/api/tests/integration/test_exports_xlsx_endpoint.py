"""Integration tests for GET /snapshots/{id}/export.xlsx."""
from __future__ import annotations

import io
from typing import Any
from uuid import uuid4

import pytest

pytestmark = pytest.mark.integration


async def _seed_snapshot(client: Any) -> str:
    ent = await client.post(
        "/entities",
        json={"name": "Test Trust", "registration_number": None, "notes": None},
    )
    eid = ent.json()["id"]
    prop = await client.post(
        "/properties",
        json={
            "entity_id": eid,
            "name": "Test Property",
            "address": "1 Test St",
            "property_type": "office",
            "notes": None,
        },
    )
    pid = prop.json()["id"]
    snap = await client.post(
        f"/properties/{pid}/snapshots",
        json={
            "valuation_date": "2026-03-15",
            "tenants": [
                {
                    "description": "T1",
                    "tenant_name": None,
                    "rentable_area_m2": "100",
                    "rent_per_m2_pm": "150",
                    "annual_escalation_pct": "0.08",
                    "next_escalation_date": None,
                    "lease_period_text": None,
                    "lease_expiry_date": None,
                }
            ],
            "parking": [],
            "monthly_operating_expenses": "10000",
            "vacancy_allowance_pct": "0.05",
            "cap_rate": "0.10",
        },
    )
    assert snap.status_code == 201, snap.text
    return str(snap.json()["id"])


async def test_get_xlsx_returns_workbook(viewer_client: Any, valuer_client: Any) -> None:
    snap_id = await _seed_snapshot(valuer_client)
    resp = await viewer_client.get(f"/snapshots/{snap_id}/export.xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in resp.headers["content-disposition"]
    assert resp.headers["cache-control"] == "public, max-age=31536000, immutable"
    assert resp.headers["etag"] == f'"{snap_id}"'
    # Round-trip: parsing the response with openpyxl should succeed.
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert "sheet1" in wb.sheetnames or wb.active is not None


async def test_get_xlsx_304_on_matching_etag(
    viewer_client: Any, valuer_client: Any
) -> None:
    snap_id = await _seed_snapshot(valuer_client)
    resp = await viewer_client.get(
        f"/snapshots/{snap_id}/export.xlsx",
        headers={"If-None-Match": f'"{snap_id}"'},
    )
    assert resp.status_code == 304


async def test_get_xlsx_404_unknown(viewer_client: Any) -> None:
    resp = await viewer_client.get(f"/snapshots/{uuid4()}/export.xlsx")
    assert resp.status_code == 404
